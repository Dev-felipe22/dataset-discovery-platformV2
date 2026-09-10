"""On-demand Excel export for one eval run — same shape as the original
hf-queries.py workbook (a sheet per query with live formulas, plus a
Summary sheet), but built from a persisted run instead of an in-memory
capture. Only called when the user explicitly asks for an export; nothing
here runs as part of every benchmark run.
"""

from __future__ import annotations

from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="2C3E50")
LEGEND_FILL = PatternFill("solid", fgColor="FFF3DC")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
HELPER_FILL = PatternFill("solid", fgColor="F2F2F2")
METRIC_FILL = PatternFill("solid", fgColor="E8EEF4")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

F_BASE = Font(name="Arial", size=10)
F_BOLD = Font(name="Arial", size=10, bold=True)
F_HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name="Arial", size=14, bold=True, color="2C3E50")
F_SMALL = Font(name="Arial", size=9, italic=True, color="666666")
F_HELPER = Font(name="Arial", size=9, color="999999")

# Visible columns
COL_RANK, COL_ID, COL_TITLE, COL_SIZE, COL_MOD, COL_DL, COL_LIKES, COL_REL, COL_NOTES = range(1, 10)
GAP_COL = 10
# Helper columns (rel flag, running count, precision@i, AP contribution, DCG, IDCG)
COL_J_REL, COL_J_CUM, COL_J_P, COL_J_AP, COL_J_DCG, COL_J_IDCG = range(11, 17)

METRIC_LABELS = ["Precision@{k}", "Recall@{k}", "F1@{k}", "MRR", "AP", "nDCG@{k}"]


def _safe_sheet_name(name: str, used: set) -> str:
    cleaned = "".join(c for c in name if c not in '[]:*?/\\') or "Query"
    base = cleaned[:31]
    candidate = base
    n = 2
    while candidate.lower() in used:
        suffix = f"_{n}"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _build_query_sheet(wb: Workbook, q: Dict[str, Any], k: int, sheet_name: str):
    ws = wb.create_sheet(sheet_name)
    results: List[Dict[str, Any]] = q["results"]
    total_relevant = q["total_relevant"]

    ws.cell(row=1, column=1, value=f"{q['query_id']} - {q.get('label', '')}").font = F_TITLE
    meta = [
        ("Query", q["query"]),
        ("Total relevant in corpus", total_relevant),
        ("Results captured", len(results)),
        ("k", k),
    ]
    for i, (lab, val) in enumerate(meta, start=2):
        ws.cell(row=i, column=1, value=lab).font = F_BOLD
        ws.cell(row=i, column=2, value=val).font = F_BASE

    legend_r = len(meta) + 3
    c = ws.cell(
        row=legend_r,
        column=1,
        value="The yellow Relevant column reflects judgments made in the app. Edit it here "
        "and the metrics below recalculate in Excel too, but it won't be written back to the app.",
    )
    c.font = Font(name="Arial", size=9, italic=True)
    c.fill = LEGEND_FILL
    ws.merge_cells(start_row=legend_r, start_column=1, end_row=legend_r, end_column=COL_NOTES)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[legend_r].height = 30

    head_r = legend_r + 2
    headers = {
        COL_RANK: "Rank", COL_ID: "Dataset ID", COL_TITLE: "Title", COL_SIZE: "Size class",
        COL_MOD: "Modalities", COL_DL: "Downloads", COL_LIKES: "Likes", COL_REL: "Relevant",
        COL_NOTES: "Notes", COL_J_REL: "rel", COL_J_CUM: "cum", COL_J_P: "P@i",
        COL_J_AP: "AP", COL_J_DCG: "DCG", COL_J_IDCG: "IDCG",
    }
    for ci in range(1, COL_J_IDCG + 1):
        h = headers.get(ci, "")
        cell = ws.cell(row=head_r, column=ci, value=h)
        if ci <= COL_NOTES:
            cell.font = F_HEAD
            cell.fill = HEADER_FILL
        else:
            cell.font = F_HELPER
            cell.fill = HELPER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    first = head_r + 1
    n = len(results)
    last = head_r + n if n else head_r
    rank_col = get_column_letter(COL_RANK)

    for i, r in enumerate(results):
        row_i = first + i
        extra = r.get("extra") or {}
        ws.cell(row=row_i, column=COL_RANK, value=r["rank"]).font = F_BASE
        ws.cell(row=row_i, column=COL_ID, value=r["dataset_id"]).font = F_BASE
        ws.cell(row=row_i, column=COL_TITLE, value=r.get("title") or "").font = F_BASE
        ws.cell(row=row_i, column=COL_SIZE, value=extra.get("size_class") or "").font = F_BASE
        mods = extra.get("modalities") or []
        ws.cell(row=row_i, column=COL_MOD, value=", ".join(mods) if isinstance(mods, list) else str(mods)).font = F_BASE
        ws.cell(row=row_i, column=COL_DL, value=extra.get("downloads")).font = F_BASE
        ws.cell(row=row_i, column=COL_LIKES, value=extra.get("likes")).font = F_BASE

        relevant = r.get("relevant")
        rel_value = "yes" if relevant is True else ("no" if relevant is False else "")
        g = ws.cell(row=row_i, column=COL_REL, value=rel_value)
        g.font = F_BOLD
        g.fill = INPUT_FILL
        g.alignment = Alignment(horizontal="center")
        ws.cell(row=row_i, column=COL_NOTES).fill = INPUT_FILL

        for ci in range(1, COL_NOTES + 1):
            cc = ws.cell(row=row_i, column=ci)
            cc.border = BORDER
            if ci in (COL_TITLE, COL_NOTES):
                cc.alignment = Alignment(vertical="center", wrap_text=True)

        rel_col = get_column_letter(COL_REL)
        ws.cell(row=row_i, column=COL_J_REL, value=f'=IF(LOWER({rel_col}{row_i})="yes",1,0)')
        ws.cell(row=row_i, column=COL_J_CUM, value=f"=SUM(${get_column_letter(COL_J_REL)}${first}:{get_column_letter(COL_J_REL)}{row_i})")
        ws.cell(row=row_i, column=COL_J_P, value=f"={get_column_letter(COL_J_CUM)}{row_i}/{rank_col}{row_i}")
        ws.cell(row=row_i, column=COL_J_AP, value=f"={get_column_letter(COL_J_REL)}{row_i}*{get_column_letter(COL_J_P)}{row_i}")
        ws.cell(row=row_i, column=COL_J_DCG, value=f"={get_column_letter(COL_J_REL)}{row_i}/LOG({rank_col}{row_i}+1,2)")
        ws.cell(
            row=row_i, column=COL_J_IDCG,
            value=f"=IF({rank_col}{row_i}<={total_relevant},1/LOG({rank_col}{row_i}+1,2),0)",
        )
        for ci in range(COL_J_REL, COL_J_IDCG + 1):
            ws.cell(row=row_i, column=ci).font = F_HELPER
            ws.cell(row=row_i, column=ci).number_format = "0.0000"

    m_r = last + 2
    ws.cell(row=m_r, column=1, value="Metrics").font = Font(name="Arial", size=12, bold=True, color="2C3E50")
    names = [lbl.format(k=k) for lbl in METRIC_LABELS]

    if n == 0:
        for i, nm in enumerate(names, start=1):
            ws.cell(row=m_r + i, column=1, value=nm).font = F_BOLD
            ws.cell(row=m_r + i, column=2, value=0).font = F_BASE
    else:
        rel_range = f"${get_column_letter(COL_REL)}${first}:${get_column_letter(COL_REL)}${last}"
        found = f'COUNTIF({rel_range},"yes")'
        ap_range = f"${get_column_letter(COL_J_AP)}${first}:${get_column_letter(COL_J_AP)}${last}"
        dcg_range = f"${get_column_letter(COL_J_DCG)}${first}:${get_column_letter(COL_J_DCG)}${last}"
        idcg_range = f"${get_column_letter(COL_J_IDCG)}${first}:${get_column_letter(COL_J_IDCG)}${last}"
        formulas = [
            f"={found}/{k}",
            f"={found}/{total_relevant}",
            f"=IFERROR(2*B{m_r+1}*B{m_r+2}/(B{m_r+1}+B{m_r+2}),0)",
            f'=IFERROR(1/MATCH("yes",{rel_range},0),0)',
            f"=IFERROR(SUM({ap_range})/{total_relevant},0)",
            f"=IFERROR(SUM({dcg_range})/SUM({idcg_range}),0)",
        ]
        for i, (nm, f) in enumerate(zip(names, formulas), start=1):
            ws.cell(row=m_r + i, column=1, value=nm).font = F_BOLD
            cell = ws.cell(row=m_r + i, column=2, value=f)
            cell.font = F_BASE
            cell.fill = METRIC_FILL
            cell.number_format = "0.0000"
            cell.border = BORDER

        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{get_column_letter(COL_REL)}{first}:{get_column_letter(COL_REL)}{last}")
        ws.freeze_panes = ws.cell(row=first, column=1)

    for col, w in zip("ABCDEFGHI", [7, 34, 32, 12, 20, 11, 8, 11, 24]):
        ws.column_dimensions[col].width = w
    for ci in range(COL_J_REL, COL_J_IDCG + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 9
    ws.column_dimensions[get_column_letter(GAP_COL)].width = 3

    return m_r


def _build_summary(wb: Workbook, run: Dict[str, Any], queries: List[Dict[str, Any]], sheet_names: Dict[str, str], metric_rows: Dict[str, int]):
    ws = wb["Summary"]
    k = run["k"]

    ws["A1"] = "Retrieval Evaluation Summary"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="2C3E50")
    created = run.get("created_at")
    created_str = created.strftime("%Y-%m-%d %H:%M") if isinstance(created, datetime) else str(created or "")
    ws["A2"] = f"Run: {run['label']}    Engine: {run['engine']}    k={k}    {created_str}"
    ws["A2"].font = F_SMALL

    head_r = 4
    headers = ["Query", "Description", f"P@{k}", f"R@{k}", f"F1@{k}", "MRR", "AP", f"nDCG@{k}"]
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=head_r, column=ci, value=h)
        cell.font = F_HEAD
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, q in enumerate(queries):
        r = head_r + 1 + i
        sheet = sheet_names[q["query_id"]]
        m_r = metric_rows[q["query_id"]]
        ws.cell(row=r, column=1, value=q["query_id"]).font = F_BOLD
        ws.cell(row=r, column=2, value=q.get("label", "")).font = F_BASE
        for j in range(6):
            cell = ws.cell(row=r, column=3 + j, value=f"='{sheet}'!B{m_r + 1 + j}")
            cell.font = F_BASE
            cell.number_format = "0.0000"
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
        for ci in (1, 2):
            ws.cell(row=r, column=ci).border = BORDER

    mean_r = head_r + 1 + len(queries)
    ws.cell(row=mean_r, column=1, value="MEAN").font = Font(name="Arial", size=10, bold=True, color="2C3E50")
    ws.cell(row=mean_r, column=2, value="across all queries (MEAN of AP = MAP)").font = F_SMALL
    for j in range(6):
        col = get_column_letter(3 + j)
        cell = ws.cell(row=mean_r, column=3 + j, value=f"=AVERAGE({col}{head_r + 1}:{col}{mean_r - 1})")
        cell.font = F_BOLD
        cell.fill = METRIC_FILL
        cell.number_format = "0.0000"
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    for ci in (1, 2):
        ws.cell(row=mean_r, column=ci).border = BORDER

    for col, w in zip("ABCDEFGH", [12, 30, 11, 11, 11, 11, 11, 11]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=head_r + 1, column=1)


def build_workbook(run: Dict[str, Any], queries: List[Dict[str, Any]]) -> Workbook:
    """run: the run summary dict from eval_service.get_run_detail()['run'].
    queries: the ['queries'] list from the same call."""
    wb = Workbook()
    wb.active.title = "Summary"

    used_names: set = set()
    sheet_names: Dict[str, str] = {}
    metric_rows: Dict[str, int] = {}
    for q in queries:
        name = _safe_sheet_name(q["query_id"], used_names)
        sheet_names[q["query_id"]] = name
        metric_rows[q["query_id"]] = _build_query_sheet(wb, q, run["k"], name)

    _build_summary(wb, run, queries, sheet_names, metric_rows)
    return wb
